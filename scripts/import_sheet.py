#!/usr/bin/env python3
"""Import the EKY Directory Google Sheet (exported as .xlsx) into data/entries/<country>/*.json.

Usage:  python3 scripts/import_sheet.py path/to/sheet.xlsx [--country el-salvador]

- One JSON file per entry. The sheet stays the place where people type; the JSON files are the
  single source of truth the site is built from (versioned in git).
- IDs are stable: an entry keeps its ID once assigned (data/registry.json). If the sheet's
  "ID (Claude fills)" column has a value, that wins.
- Curated extras for EKY clients live in data/profiles/<id>.json and are never touched here.
- Rows with Status "Remove" are deleted from data/entries (the page disappears on next build).
"""
import json, re, sys, unicodedata, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CATS = json.loads((ROOT / "data/categories.json").read_text())["categories"]
TAB_TO_CAT = {c["sheetTab"]: c for c in CATS}

def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return s or "entry"

def clean(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()

def split_list(s):
    return [x.strip() for x in re.split(r"[,;\n]", s) if x.strip()]

def yesno(s):
    s = s.lower()
    return True if s in ("yes", "sí", "si", "true") else False if s in ("no", "false") else None

def num(s):
    try:
        return float(s.replace(",", "."))
    except Exception:
        return None

def url(s):
    s = s.strip()
    if not s:
        return ""
    if not re.match(r"^https?://", s):
        s = "https://" + s
    return s

def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__); sys.exit(1)
    xlsx = Path(args[0])
    country = "el-salvador"
    if "--country" in args:
        country = args[args.index("--country") + 1]
    out_dir = ROOT / "data/entries" / country
    out_dir.mkdir(parents=True, exist_ok=True)
    reg_path = ROOT / "data/registry.json"
    registry = json.loads(reg_path.read_text()) if reg_path.exists() else {}
    used = set(registry.values())

    wb = load_workbook(xlsx, data_only=True)
    seen, removed, problems = set(), [], []
    counts = {}
    for tab, cat in TAB_TO_CAT.items():
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        headers = [clean(c.value) for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {h: clean(v) for h, v in zip(headers, r) if h}
            name = row.get("Name", "")
            if not name:
                continue
            status = row.get("Status", "") or "To review"
            key = f"{cat['key']}::{slugify(name)}"
            eid = row.get("ID (Claude fills)") or registry.get(key)
            if not eid:
                base = slugify(name)
                eid, n = base, 2
                while eid in used:
                    eid = f"{base}-{n}"; n += 1
            registry[key] = eid
            used.add(eid)
            path = out_dir / f"{eid}.json"
            if status == "Remove":
                if path.exists():
                    path.unlink(); removed.append(eid)
                continue
            seen.add(eid)
            e = build_entry(eid, country, cat, row, status)
            for p in validate(e):
                problems.append(f"{tab} / {name}: {p}")
            path.write_text(json.dumps(e, ensure_ascii=False, indent=2) + "\n")
            counts[tab] = counts.get(tab, 0) + 1

    # entries that vanished from the sheet are kept (never delete silently) but reported
    orphans = [p.stem for p in out_dir.glob("*.json") if p.stem not in seen]
    reg_path.write_text(json.dumps(registry, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
    print("Imported:", counts)
    if removed: print("Removed:", removed)
    if orphans: print("In data but not in sheet (kept):", orphans)
    if problems:
        print(f"\n{len(problems)} things to check:")
        for p in problems: print(" -", p)

def build_entry(eid, country, cat, row, status):
    g = row.get
    kind = cat["kind"]
    e = {
        "id": eid, "kind": kind, "country": country, "category": cat["key"],
        "status": status, "ekyClient": yesno(g("EKY client", "")) is True,
        "name": g("Name"),
    }
    if kind == "organization":
        e["type"] = "Startup" if cat["key"] == "startups" else g("Type", "")
        e["tags"] = split_list(g("Tags", "") or g("Sector tags", ""))
        also = g("Also listed in", "")
        also_cat = TAB_TO_CAT.get(also)
        e["alsoIn"] = [also_cat["key"]] if also_cat else []
        e["description"] = {"en": g("Description (EN)", "") or g("One-liner (EN)", ""),
                             "es": g("Descripción (ES)", "") or g("Una línea (ES)", "")}
        e["website"] = url(g("Official website", ""))
        e["links"] = {k: url(g(h, "")) for k, h in (("instagram", "Instagram"), ("facebook", "Facebook"),
                      ("linkedin", "LinkedIn"), ("tiktok", "TikTok"), ("youtube", "YouTube"), ("x", "X"))}
        if g("LinkedIn (company)"):
            e["links"]["linkedin"] = url(g("LinkedIn (company)"))
        e["maps"] = url(g("Google Maps link", ""))
        lat, lng = num(g("Latitude", "")), num(g("Longitude", ""))
        e["geo"] = {"lat": lat, "lng": lng} if lat is not None and lng is not None else None
        e["location"] = {"department": g("Department", ""), "municipality": g("Municipality", ""),
                         "district": g("District", ""), "address": g("Address", ""), "city": g("City", "")}
        e["phone"] = g("Public phone", "")
        e["email"] = g("Public email", "")
        e["languages"] = split_list(g("Languages", ""))
        e["acceptsBitcoin"] = yesno(g("Accepts Bitcoin", ""))
        extra = {}
        for h, k in (("Price range", "priceRange"), ("Booking link", "bookingUrl"), ("RNT registered", "rnt"),
                     ("Opening hours", "hours"), ("Menu link", "menuUrl"), ("Reservations link", "reservationsUrl"),
                     ("Departments served", "areaServed"), ("Founded year", "foundingYear"), ("Memberships", "memberships"),
                     ("Zones served", "zonesServed"), ("Helps foreign or diaspora buyers", "foreignBuyers"),
                     ("Public or private", "publicPrivate"), ("Curriculum", "curriculum"), ("Admissions link", "admissionsUrl"),
                     ("Accreditation", "accreditation"), ("Other campuses", "otherCampuses"),
                     ("Markets served", "marketsServed"), ("Team size", "teamSize"), ("Main products", "products"),
                     ("Exports", "exports"), ("Export markets", "exportMarkets"), ("Certifications", "certifications"),
                     ("Stage", "stage"), ("Founders / leaders", "founders"), ("Programs / investors", "programs"),
                     ("Hiring", "hiring")):
            v = g(h, "")
            if not v:
                continue
            if k in ("rnt", "foreignBuyers", "exports", "hiring"):
                v = yesno(v)
            elif k in ("bookingUrl", "menuUrl", "reservationsUrl", "admissionsUrl"):
                v = url(v)
            elif k in ("founders", "programs", "areaServed", "marketsServed", "exportMarkets", "certifications", "memberships"):
                v = split_list(v)
            extra[k] = v
        e["extra"] = extra
    else:
        e["type"] = "Person"
        e["tags"] = split_list(g("Topic tags", "") or g("Niche tags", ""))
        e["role"] = g("Role / headline", "")
        e["ventures"] = split_list(g("Ventures", ""))
        e["openTo"] = split_list(g("Open to", ""))
        e["description"] = {"en": g("Short bio (EN)", ""), "es": g("Bio corta (ES)", "")}
        e["website"] = url(g("Personal website", ""))
        e["links"] = {k: url(g(h, "")) for k, h in (("linkedin", "LinkedIn"), ("instagram", "Instagram"), ("x", "X"),
                      ("youtube", "YouTube"), ("tiktok", "TikTok"), ("facebook", "Facebook"))}
        e["geo"] = None  # people never get coordinates
        e["location"] = {"city": g("City", ""), "department": g("Department", "")}
        e["languages"] = split_list(g("Languages", ""))
        e["email"] = g("Public business email", "")
        extra = {}
        if g("Main platform"): extra["mainPlatform"] = g("Main platform")
        if g("Audience size"): extra["audience"] = g("Audience size")
        if g("Audience checked on"): extra["audienceCheckedOn"] = g("Audience checked on")
        e["extra"] = extra
    e["sources"] = [url(s) if s.startswith(("http", "www")) else s for s in split_list(g("Source URL", ""))]
    e["lastChecked"] = g("Last checked", "")
    # "Internal notes" from the sheet are deliberately NOT copied: this repository is public.
    return e

def validate(e):
    out = []
    if not e["sources"]:
        out.append("no source")
    if e["kind"] == "organization" and not e.get("type"):
        out.append("no type")
    if e.get("geo"):
        lat, lng = e["geo"]["lat"], e["geo"]["lng"]
        if not (13.1 <= lat <= 14.5 and -90.2 <= lng <= -87.6):
            out.append(f"coordinates {lat},{lng} are outside El Salvador (swapped?)")
    for k, v in (e.get("links") or {}).items():
        if v and k == "linkedin" and "linkedin.com" not in v:
            out.append("LinkedIn link is not a linkedin.com URL")
        if v and ("google.com/search" in v or "/search?" in v):
            out.append(f"{k} is a search link, not a profile")
    if e.get("phone") and re.sub(r"\D", "", e["phone"])[:1] in ("6", "7") and e["kind"] == "person":
        out.append("person has a phone number")
    return out

if __name__ == "__main__":
    main()
